from __future__ import annotations

import csv as _csv
import re
from collections.abc import Iterable, Iterator, Sequence
from datetime import date
from typing import Any

from openpyxl import load_workbook

from sheetrender.column_keys import sanitize_columns
from sheetrender.templating import date_needs_confirmation, parse_date


class TooManyCellsError(ValueError):
    pass


DEFAULT_MAX_CELLS = 2_000_000
_SAMPLE_ROW_COUNT = 5

MAX_XLSX_CELL_CHARACTERS = 32_767


class CellTooLongError(ValueError):
    pass


def _too_many_cells_message(max_cells: int) -> str:
    return (
        f"This file has too many populated cells (max {max_cells:,}). "
        "Split it into smaller sheets and try again."
    )


# Only plain decimal notation becomes a number. Anything float() or int() would
# also accept stays text, because it changes what the cell says: "02134" is a
# zip code, "Nan" and "Infinity" are names, "1_000" and "1e3" are codes.
_CSV_INT_RE = re.compile(r"-?(0|[1-9][0-9]*)")
_CSV_FLOAT_RE = re.compile(r"-?(0|[1-9][0-9]*)\.[0-9]+")


def _convert_csv_cell(cell: str) -> int | float | str | None:
    cell = cell.strip()
    if cell == "":
        return None
    if _CSV_INT_RE.fullmatch(cell):
        return int(cell)
    if _CSV_FLOAT_RE.fullmatch(cell):
        return float(cell)
    return cell


def _fit_row(row: Sequence[Any], width: int) -> tuple[Any, ...]:
    """Trim or None-pad a row to exactly `width` values.

    Padding matters: templates run under StrictUndefined, so a short row that
    left its trailing keys out of the row dict would fail to render.
    """
    return tuple(row[:width]) + (None,) * (width - len(row))


def _is_empty_row(values: tuple[Any, ...]) -> bool:
    return all(v is None for v in values)


def _json_row(keys: list[str], values: tuple[Any, ...]) -> dict[str, Any]:
    # Only preview samples cross the JSON boundary; iter_rows preserves cells.
    return {
        key: value.isoformat() if isinstance(value, date) else value
        for key, value in zip(keys, values, strict=True)
    }


def _xlsx_data_rows(rows: Iterable[Sequence[Any]], width: int) -> Iterator[tuple[Any, ...]]:
    for row in rows:
        values = _fit_row(row, width)
        if not _is_empty_row(values):
            yield values


def _csv_data_rows(reader: Iterable[list[str]], width: int) -> Iterator[tuple[Any, ...]]:
    for row in reader:
        values = _fit_row([_convert_csv_cell(cell) for cell in row], width)
        if not _is_empty_row(values):
            yield values


def _summarize(
    sheet_name: str,
    originals: list[str],
    data_rows: Iterator[tuple[Any, ...]],
    max_cells: int,
) -> dict[str, Any]:
    """Build the dataset description shared by parse_xlsx and parse_csv.

    `data_rows` yields non-empty rows already fitted to len(originals).
    """
    keys = sanitize_columns(originals)
    # Counts populated cells, header included, so sparse files aren't penalized
    # for their width.
    populated_cells = sum(1 for original in originals if original)
    if populated_cells > max_cells:
        raise TooManyCellsError(_too_many_cells_message(max_cells))

    # Infer within the same bounded window as the app's compatibility shim.
    inferred_types = ["string"] * len(keys)
    pending = set(range(len(keys)))
    candidates: set[int] = set()
    # Ambiguous forms (1.2.2024) are also version numbers, so a column typed
    # from them stays under watch: any later non-date cell makes it a string.
    watched: set[int] = set()
    sample_rows: list[dict[str, Any]] = []
    row_count = 0
    for values in data_rows:
        row_count += 1
        populated_cells += sum(1 for v in values if v is not None)
        if populated_cells > max_cells:
            raise TooManyCellsError(_too_many_cells_message(max_cells))
        if row_count <= 5000:
            for index in list(pending):
                value = values[index]
                if value is None or (isinstance(value, str) and not value.strip()):
                    continue
                if parse_date(value, excel_serial=False) is not None:
                    if index in watched:
                        continue
                    if date_needs_confirmation(value):
                        if index not in candidates:
                            candidates.add(index)
                            continue
                        inferred_types[index] = "date"
                        watched.add(index)
                        continue
                    inferred_types[index] = "date"
                elif index in watched:
                    inferred_types[index] = "string"
                elif index not in candidates and isinstance(value, (int, float)) and not isinstance(value, bool):
                    inferred_types[index] = "number"
                pending.remove(index)
        if len(sample_rows) < _SAMPLE_ROW_COUNT:
            sample_rows.append(_json_row(keys, values))

    columns = [
        {"original": original, "key": key, "inferred_type": inferred_type}
        for original, key, inferred_type in zip(originals, keys, inferred_types, strict=True)
    ]
    return {
        "sheet_name": sheet_name,
        "columns": columns,
        "row_count": row_count,
        "sample_rows": sample_rows,
    }


def parse_xlsx(path: str, max_cells: int | None = None) -> dict[str, Any]:
    if max_cells is None:
        max_cells = DEFAULT_MAX_CELLS
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header_row = next(rows, None)
        if header_row is None:
            raise ValueError("Workbook has no header row")
        originals = ["" if h is None else str(h) for h in header_row]
        return _summarize(ws.title, originals, _xlsx_data_rows(rows, len(originals)), max_cells)
    finally:
        wb.close()


def parse_csv(path: str, max_cells: int | None = None) -> dict[str, Any]:
    if max_cells is None:
        max_cells = DEFAULT_MAX_CELLS
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = _csv.reader(f)
        header_row = next(reader, None)
        if header_row is None:
            raise ValueError("CSV has no header row")
        originals = [h.strip() for h in header_row]
        return _summarize("csv", originals, _csv_data_rows(reader, len(originals)), max_cells)


def _iter_rows_csv(path: str, keys: list[str]) -> Iterator[dict[str, Any]]:
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = _csv.reader(f)
        next(reader, None)
        for values in _csv_data_rows(reader, len(keys)):
            yield dict(zip(keys, values, strict=True))


def _iter_rows_xlsx(path: str, keys: list[str]) -> Iterator[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb.worksheets[0].iter_rows(values_only=True)
        next(rows, None)
        for values in _xlsx_data_rows(rows, len(keys)):
            yield dict(zip(keys, values, strict=True))
    finally:
        wb.close()


def iter_rows(path: str, columns: list[dict[str, Any]]) -> Iterator[dict[str, Any]]:
    """Yield each non-empty data row as a dict keyed by the columns' template keys."""
    keys = [c["key"] for c in columns]
    if path.lower().endswith(".csv"):
        return _iter_rows_csv(path, keys)
    return _iter_rows_xlsx(path, keys)
