from pathlib import Path

import pytest

from sheetrender.sheets import parse_xlsx

# The backend suite read this workbook from the repo's examples/ directory. It
# travels with the tests here instead, so the package's suite is self-contained.
XLSX = Path(__file__).parent / "data" / "performance-reports.xlsx"


def test_parse_xlsx():
    result = parse_xlsx(str(XLSX))
    assert len(result["columns"]) == 16
    keys = [c["key"] for c in result["columns"]]
    assert "employee_name" in keys
    assert "performance_pct" in keys
    assert "bonus_eligible" in keys
    assert 90 <= result["row_count"] <= 110
    assert isinstance(result["sample_rows"][0]["sales_target"], (int, float))


def test_xlsx_dates_are_detected_and_json_serializable(tmp_path):
    import json
    from datetime import date, datetime

    from openpyxl import Workbook

    from sheetrender.sheets import iter_rows

    path = tmp_path / "dates.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Native", "Timestamp", "ISO", "Amount", "Serial", "Boolean"])
    sheet.append([date(2026, 9, 22), datetime(2026, 9, 22, 14, 30), "2026-09-22", 12.5, 46287, True])
    workbook.save(path)
    result = parse_xlsx(str(path))
    assert [column["inferred_type"] for column in result["columns"]] == [
        "date", "date", "date", "number", "number", "string",
    ]
    rows = list(iter_rows(str(path), result["columns"]))
    assert rows[0]["native"] == datetime(2026, 9, 22)
    assert rows[0]["timestamp"] == datetime(2026, 9, 22, 14, 30)
    assert str(rows[0]["timestamp"]) == "2026-09-22 14:30:00"
    assert result["sample_rows"][0]["timestamp"] == "2026-09-22T14:30:00"
    assert rows[0]["iso"] == "2026-09-22"
    json.dumps(result)

    from sheetrender.grouping import grouped_render_units
    from sheetrender.templating import render_row

    grouped = grouped_render_units(str(path), result["columns"], None, {"group_by": ["iso"]})[0]
    assert grouped.rows[0] == rows[0]
    assert render_row("{{ timestamp }}", rows[0]) == "2026-09-22 14:30:00"
    assert render_row('{{ timestamp | date("%Y-%m-%d") }}', rows[0]) == "2026-09-22"


def test_csv_dates_use_first_populated_cell_including_after_sample_window(tmp_path):
    from sheetrender.sheets import iter_rows, parse_csv

    path = tmp_path / "dates.csv"
    path.write_text(
        "ISO,US,European,Late,Code,Number,Mixed,Invalid\n"
        "2026-09-22,09/22/2026,22.09.2026,,02134,46287,hello,2026-02-30\n"
        + "2026-09-23,,23.09.2026,,,,,,\n" * 5
        + "2026-09-24,,,2026-09-25,46287,7,2026-09-22,none\n"
    )
    result = parse_csv(str(path))
    by_key = {column["key"]: column["inferred_type"] for column in result["columns"]}
    assert by_key == {
        "iso": "date", "us": "date", "european": "date", "late": "date",
        "code": "string", "number": "number", "mixed": "string", "invalid": "string",
    }
    assert result["row_count"] == 7
    first = next(iter_rows(str(path), result["columns"]))
    assert first["us"] == "09/22/2026"
    assert first["code"] == "02134"
    assert first["number"] == 46287


def test_date_ingest_keeps_the_cell_cap(tmp_path):
    import pytest

    from sheetrender.sheets import TooManyCellsError, parse_csv

    path = tmp_path / "dates.csv"
    path.write_text("Issued\n2026-09-22\n2026-09-23\n")
    with pytest.raises(TooManyCellsError):
        parse_csv(str(path), max_cells=2)


@pytest.mark.parametrize("value", [
    "1.2.2024", "10.11.2023", "1/2/2024", "1-2-2024", "555-123-4567",
    "02134", "12345", "20240922", "1899-12-31", "2101-01-01", "02/30/2024",
])
def test_single_ambiguous_date_or_identifier_stays_non_date(tmp_path, value):
    from sheetrender.sheets import parse_csv

    path = tmp_path / "values.csv"
    path.write_text("Value\n" + value + "\n")
    assert parse_csv(str(path))["columns"][0]["inferred_type"] != "date"


@pytest.mark.parametrize("first,second", [
    ("1.2.2024", "2.3.2024"), ("10.11.2023", "11.12.2023"),
    ("1/2/2024", "2/3/2024"), ("1-2-2024", "2-3-2024"),
])
def test_ambiguous_dates_need_two_nonblank_agreeing_cells(tmp_path, first, second):
    from sheetrender.sheets import parse_csv

    path = tmp_path / "dates.csv"
    path.write_text("Value,Other\n" + first + ",a\n ,b\n" + second + ",c\n")
    assert parse_csv(str(path))["columns"][0]["inferred_type"] == "date"
    path.write_text("Value\n" + first + "\nversion\n" + second + "\n")
    assert parse_csv(str(path))["columns"][0]["inferred_type"] == "string"


def test_ambiguous_date_column_with_a_later_version_string_stays_string(tmp_path):
    from sheetrender.sheets import parse_csv

    path = tmp_path / "versions.csv"
    path.write_text("Version,Due\n1.2.2024,1.2.2024\n10.11.2023,10.11.2023\n2.0.1,\n")
    by_key = {column["key"]: column["inferred_type"] for column in parse_csv(str(path))["columns"]}
    assert by_key == {"version": "string", "due": "date"}
    # Confirmation by an unambiguous date still leaves the column under watch.
    path.write_text("Value,Due\n1.2.2024,1.2.2024\n2024-03-02,2024-03-02\nversion,\n")
    by_key = {column["key"]: column["inferred_type"] for column in parse_csv(str(path))["columns"]}
    assert by_key == {"value": "string", "due": "date"}


def test_date_detection_stops_at_5000_rows(tmp_path):
    from sheetrender.sheets import parse_csv

    path = tmp_path / "bounded.csv"
    path.write_text(
        "Row,Within,Late,Unconfirmed,Blank\n"
        + "".join(f"{i},,,,\n" for i in range(1, 5000))
        + "5000,2026-09-22,,1.2.2024,\n"
        + "5001,,2026-09-22,2.3.2024,\n"
    )
    result = parse_csv(str(path))
    assert result["row_count"] == 5001
    assert [column["inferred_type"] for column in result["columns"]] == [
        "number", "date", "string", "string", "string",
    ]


def test_time_and_duration_cells_keep_sample_rows_json(tmp_path):
    """openpyxl reads 09:30 as datetime.time and [h]:mm as timedelta; either one
    in the first rows made the dataset's sample rows unserialisable."""
    import json
    from datetime import time, timedelta

    from openpyxl import Workbook

    from sheetrender.sheets import iter_rows
    from sheetrender.templating import render_row

    path = tmp_path / "appointments.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Client", "Appointment time", "Hours"])
    sheet.append(["Ann", time(9, 30), timedelta(hours=1, minutes=30)])
    workbook.save(path)

    result = parse_xlsx(str(path))

    json.dumps(result)
    assert result["sample_rows"][0] == {
        "client": "Ann", "appointment_time": "09:30:00", "hours": "1:30:00",
    }
    row = next(iter_rows(str(path), result["columns"]))
    template = "{{ appointment_time }} {{ hours }}"
    assert render_row(template, row) == render_row(template, result["sample_rows"][0])


def test_zero_padded_number_formats_keep_their_leading_zeros(tmp_path):
    """Excel's "Zip Code" format stores 2134 and shows 02134; the document used
    to print 2134. Formats of zeros and literal dashes are applied at read."""
    from openpyxl import Workbook

    from sheetrender.sheets import iter_rows

    path = tmp_path / "addresses.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Name", "ZIP", "ZIP+4", "SSN", "Badge", "Count", "Price", "Wide"])
    sheet.append(["Ann", 2134, 21340001, 12345678, 42, 7, 5, 1234567])
    for cell, fmt in zip(
        ("B2", "C2", "D2", "E2", "F2", "G2", "H2"),
        ("00000", "00000\\-0000", "000\\-00\\-0000", "000000", "0", '"$"#,##0.00', "00000"),
        strict=True,
    ):
        sheet[cell].number_format = fmt
    workbook.save(path)

    result = parse_xlsx(str(path))
    row = next(iter_rows(str(path), result["columns"]))

    assert row == {
        "name": "Ann", "zip": "02134", "zip_4": "02134-0001", "ssn": "012-34-5678",
        "badge": "000042", "count": 7, "price": 5, "wide": "1234567",
    }
    assert result["sample_rows"][0] == row
    types = {c["key"]: c["inferred_type"] for c in result["columns"]}
    assert types["zip"] == "string" and types["count"] == "number"


def test_a_hidden_first_sheet_is_skipped_for_the_first_visible_one(tmp_path):
    """A workbook whose first tab is a hidden lookup table imported that table:
    the user saw their customer list in Excel and got two rows of codes."""
    from openpyxl import Workbook

    from sheetrender.sheets import iter_rows

    path = tmp_path / "hidden.xlsx"
    workbook = Workbook()
    lookup = workbook.active
    lookup.title = "Lookup"
    lookup.append(["code", "label"])
    lookup.append([1, "one"])
    lookup.sheet_state = "hidden"
    customers = workbook.create_sheet("Customers")
    customers.append(["Name", "City"])
    customers.append(["Ann", "Oslo"])
    customers.append(["Bob", "Rome"])
    workbook.active = 1
    workbook.save(path)

    result = parse_xlsx(str(path))

    assert result["sheet_name"] == "Customers"
    assert [c["key"] for c in result["columns"]] == ["name", "city"]
    assert list(iter_rows(str(path), result["columns"])) == [
        {"name": "Ann", "city": "Oslo"}, {"name": "Bob", "city": "Rome"},
    ]
