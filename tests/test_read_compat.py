"""Stored datasets keep reading exactly as sheetrender 0.4.0 read them.

A service stores `parse_*`'s columns once, at ingest, and calls `iter_rows`
again every time it renders. Thousands of datasets were stored by 0.4.0, so
`iter_rows` without `read_options` must give back the same sheet, the same
header row, the same keys and the same value types 0.4.0 gave for them, even
where newer reading rules (hidden tabs, zero-padded formats) would answer
differently for a fresh upload.

Every LEGACY entry below is what the released 0.4.0 wheel produced for the
file its builder writes: `columns` from its `parse_*`, `rows` from its
`iter_rows`, with each value's type recorded next to it. The file itself runs
unchanged against 0.4.0 (`PYTHONPATH=<0.4.0 src> pytest
tests/test_read_compat.py -k legacy`), which is how the table was checked.
"""
from __future__ import annotations

import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from sheetrender.sheets import iter_rows


def _hidden_first_same_headers(path: Path) -> None:
    workbook = Workbook()
    hidden = workbook.active
    hidden.title = "Old"
    hidden.append(["Name", "Amount"])
    hidden.append(["Ann", 100])
    hidden.sheet_state = "hidden"
    visible = workbook.create_sheet("Current")
    visible.append(["Name", "Amount"])
    visible.append(["Bob", 900])
    workbook.active = 1
    workbook.save(path)


def _hidden_first_other_headers(path: Path) -> None:
    workbook = Workbook()
    lookup = workbook.active
    lookup.title = "Lookup"
    lookup.append(["code", "label"])
    lookup.append([1, "one"])
    lookup.sheet_state = "hidden"
    customers = workbook.create_sheet("Customers")
    customers.append(["Name", "City"])
    customers.append(["Ann", "Oslo"])
    workbook.active = 1
    workbook.save(path)


def _all_sheets_hidden_but_one_later(path: Path) -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "A"
    first.append(["x"])
    first.append([1])
    second = workbook.create_sheet("B")
    second.append(["y"])
    second.append([2])
    first.sheet_state = "veryHidden"
    workbook.active = 1
    workbook.save(path)


def _two_visible_sheets(path: Path) -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "Data"
    first.append(["Name", "Amount"])
    first.append(["Ann", 5])
    other = workbook.create_sheet("Notes")
    other.append(["Note"])
    other.append(["ignored"])
    workbook.save(path)


def _table_starts_lower_down(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A2"] = "   "
    sheet["B3"], sheet["C3"] = "Name", "Amount"
    sheet["B4"], sheet["C4"] = "Ann", 5
    sheet["B5"], sheet["C5"] = "Bob", 6
    workbook.save(path)


def _blank_first_row_then_data(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A2"], sheet["B2"] = "Ann", 5
    sheet["A3"], sheet["B3"] = "Bob", 6
    workbook.save(path)


def _zero_padded_formats(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Name", "ZIP", "Qty", "ZIP+4", "SSN", "Literal", "Wide", "Neg", "Price"])
    sheet.append(["Ann", 2134, 12, 21340001, 12345678, 5, 1234567, -5, 5])
    sheet.append(["Bob", 90210, 7, 902101234, 123456789, 6, 12, 3, 6])
    formats = ("00000", "00", "00000\\-0000", "000\\-00\\-0000", '0"0"', "00000", "00000", '"$"#,##0.00')
    for row in (2, 3):
        for column, fmt in zip("BCDEFGHI", formats, strict=True):
            sheet[f"{column}{row}"].number_format = fmt
    workbook.save(path)


def _trailing_nameless_columns(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Name", None, "Email", None, None, None])
    sheet.append(["Ann", "kept: middle", "ann@example.com", None, "kept: data", None])
    sheet.append(["Bob", None, "bob@example.com"])
    workbook.save(path)


def _cell_types(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Day", "Stamp", "Clock", "Hours", "Float", "Flag", "Text", "Int", "Formula"])
    sheet.append([
        datetime.date(2026, 9, 22), datetime.datetime(2026, 9, 22, 14, 30, 5),
        datetime.time(9, 30), datetime.timedelta(hours=25, minutes=30),
        12.5, True, "  padded  ", 7, "=1+1",
    ])
    sheet.append([None, None, None, None, None, None, None, None, None])
    sheet.append(["2026-09-22", "9/26/2026 15:59:00", None, None, 0.1, False, "02134", 0, None])
    workbook.save(path)


def _csv(text: str):
    def build(path: Path) -> None:
        path.write_bytes(text.encode("utf-8"))
    return build


# name -> (suffix, builder)
CASES = {
    "hidden_first_same_headers": (".xlsx", _hidden_first_same_headers),
    "hidden_first_other_headers": (".xlsx", _hidden_first_other_headers),
    "very_hidden_first": (".xlsx", _all_sheets_hidden_but_one_later),
    "two_visible_sheets": (".xlsx", _two_visible_sheets),
    "table_starts_lower_down": (".xlsx", _table_starts_lower_down),
    "blank_first_row_then_data": (".xlsx", _blank_first_row_then_data),
    "zero_padded_formats": (".xlsx", _zero_padded_formats),
    "trailing_nameless_columns": (".xlsx", _trailing_nameless_columns),
    "cell_types": (".xlsx", _cell_types),
    "csv_blank_header_then_data": (".csv", _csv(",\nAnn,5\nBob,6\n")),
    "csv_blank_lines_above_header": (".csv", _csv("\n,,\nName,Amount\nAnn,5\n")),
    "csv_trailing_commas": (".csv", _csv("Name,Amount,,\r\nAnn,1,,\r\nBob,2,,\r\n")),
    "csv_values": (".csv", _csv(
        "﻿Zip,Amount,Note\n02134,1.50,\"two\nlines\"\n,,\n90210,1e3, x \n-0,-12,Nan\n"
    )),
}


def build_case(tmp_path: Path, name: str) -> str:
    suffix, builder = CASES[name]
    path = tmp_path / f"{name}{suffix}"
    builder(path)
    return str(path)


def typed(rows):
    return [{key: (type(value).__name__, value) for key, value in row.items()} for row in rows]


def _cols(*pairs):
    return [{"original": original, "key": key, "inferred_type": kind} for original, key, kind in pairs]


def _t(**values):
    return values


# Generated by running each builder through the sheetrender 0.4.0 wheel from
# PyPI (identical to the v0.4.0 tag): columns from parse_*, rows from iter_rows,
# as (type name, value).
LEGACY = {
    "hidden_first_same_headers": {
        "columns": _cols(("Name", "name", "string"), ("Amount", "amount", "number")),
        "rows": [_t(name=("str", "Ann"), amount=("int", 100))],
    },
    "hidden_first_other_headers": {
        "columns": _cols(("code", "code", "number"), ("label", "label", "string")),
        "rows": [_t(code=("int", 1), label=("str", "one"))],
    },
    "very_hidden_first": {
        "columns": _cols(("x", "x", "number")),
        "rows": [_t(x=("int", 1))],
    },
    "two_visible_sheets": {
        "columns": _cols(("Name", "name", "string"), ("Amount", "amount", "number")),
        "rows": [_t(name=("str", "Ann"), amount=("int", 5))],
    },
    "table_starts_lower_down": {
        "columns": _cols(("", "col_", "string"), ("", "col__2", "string"), ("", "col__3", "string")),
        "rows": [
            _t(col_=("str", "   "), col__2=("NoneType", None), col__3=("NoneType", None)),
            _t(col_=("NoneType", None), col__2=("str", "Name"), col__3=("str", "Amount")),
            _t(col_=("NoneType", None), col__2=("str", "Ann"), col__3=("int", 5)),
            _t(col_=("NoneType", None), col__2=("str", "Bob"), col__3=("int", 6)),
        ],
    },
    "blank_first_row_then_data": {
        "columns": _cols(("", "col_", "string"), ("", "col__2", "number")),
        "rows": [
            _t(col_=("str", "Ann"), col__2=("int", 5)),
            _t(col_=("str", "Bob"), col__2=("int", 6)),
        ],
    },
    "zero_padded_formats": {
        "columns": _cols(
            ("Name", "name", "string"), ("ZIP", "zip", "number"), ("Qty", "qty", "number"),
            ("ZIP+4", "zip_4", "number"), ("SSN", "ssn", "number"), ("Literal", "literal", "number"),
            ("Wide", "wide", "number"), ("Neg", "neg", "number"), ("Price", "price", "number"),
        ),
        "rows": [
            _t(name=("str", "Ann"), zip=("int", 2134), qty=("int", 12), zip_4=("int", 21340001),
               ssn=("int", 12345678), literal=("int", 5), wide=("int", 1234567), neg=("int", -5),
               price=("int", 5)),
            _t(name=("str", "Bob"), zip=("int", 90210), qty=("int", 7), zip_4=("int", 902101234),
               ssn=("int", 123456789), literal=("int", 6), wide=("int", 12), neg=("int", 3),
               price=("int", 6)),
        ],
    },
    "trailing_nameless_columns": {
        "columns": _cols(
            ("Name", "name", "string"), ("", "col_", "string"), ("Email", "email", "string"),
            ("", "col__2", "string"), ("", "col__3", "string"), ("", "col__4", "string"),
        ),
        "rows": [
            _t(name=("str", "Ann"), col_=("str", "kept: middle"), email=("str", "ann@example.com"),
               col__2=("NoneType", None), col__3=("str", "kept: data"), col__4=("NoneType", None)),
            _t(name=("str", "Bob"), col_=("NoneType", None), email=("str", "bob@example.com"),
               col__2=("NoneType", None), col__3=("NoneType", None), col__4=("NoneType", None)),
        ],
    },
    "cell_types": {
        "columns": _cols(
            ("Day", "day", "date"), ("Stamp", "stamp", "date"), ("Clock", "clock", "string"),
            ("Hours", "hours", "string"), ("Float", "float", "number"), ("Flag", "flag", "string"),
            ("Text", "text", "string"), ("Int", "int", "number"), ("Formula", "formula", "string"),
        ),
        "rows": [
            _t(day=("datetime", datetime.datetime(2026, 9, 22, 0, 0)),
               stamp=("datetime", datetime.datetime(2026, 9, 22, 14, 30, 5)),
               clock=("time", datetime.time(9, 30)),
               hours=("timedelta", datetime.timedelta(days=1, seconds=5400)),
               float=("float", 12.5), flag=("bool", True), text=("str", "  padded  "),
               int=("int", 7), formula=("NoneType", None)),
            _t(day=("str", "2026-09-22"), stamp=("str", "9/26/2026 15:59:00"),
               clock=("NoneType", None), hours=("NoneType", None), float=("float", 0.1),
               flag=("bool", False), text=("str", "02134"), int=("int", 0),
               formula=("NoneType", None)),
        ],
    },
    "csv_blank_header_then_data": {
        "columns": _cols(("", "col_", "string"), ("", "col__2", "number")),
        "rows": [
            _t(col_=("str", "Ann"), col__2=("int", 5)),
            _t(col_=("str", "Bob"), col__2=("int", 6)),
        ],
    },
    "csv_blank_lines_above_header": {
        "columns": [],
        "rows": [],
    },
    "csv_trailing_commas": {
        "columns": _cols(
            ("Name", "name", "string"), ("Amount", "amount", "number"),
            ("", "col_", "string"), ("", "col__2", "string"),
        ),
        "rows": [
            _t(name=("str", "Ann"), amount=("int", 1), col_=("NoneType", None), col__2=("NoneType", None)),
            _t(name=("str", "Bob"), amount=("int", 2), col_=("NoneType", None), col__2=("NoneType", None)),
        ],
    },
    "csv_values": {
        "columns": _cols(("Zip", "zip", "string"), ("Amount", "amount", "number"), ("Note", "note", "string")),
        "rows": [
            _t(zip=("str", "02134"), amount=("float", 1.5), note=("str", "two\nlines")),
            _t(zip=("int", 90210), amount=("str", "1e3"), note=("str", "x")),
            _t(zip=("int", 0), amount=("int", -12), note=("str", "Nan")),
        ],
    },
}


def test_every_case_has_a_legacy_expectation():
    assert set(LEGACY) == set(CASES)


@pytest.mark.parametrize("name", sorted(CASES))
def test_legacy_columns_read_exactly_as_0_4_0(tmp_path, name):
    """No read_options: the dataset was stored by 0.4.0 (or by a caller that
    never kept them), so the file reads the way 0.4.0 read it."""
    path = build_case(tmp_path, name)
    expected = LEGACY[name]
    assert typed(iter_rows(path, expected["columns"])) == expected["rows"]


# --- Datasets ingested by this version ---------------------------------------
#
# parse_* returns read_options: how this file was read (the tab, the header
# row, the columns shown as zero-padded text). Stored next to the columns and
# passed back to iter_rows, they make every later read agree with the ingest.


def _parse(path: str) -> dict:
    from sheetrender.sheets import parse_csv, parse_xlsx

    return parse_csv(path) if path.endswith(".csv") else parse_xlsx(path)


def _read(path: str, parsed: dict) -> list[dict]:
    return list(iter_rows(path, parsed["columns"], read_options=parsed["read_options"]))


def test_a_new_ingest_reads_the_visible_tab_and_keeps_reading_it(tmp_path):
    """Both tabs share their headings, so no header matching can tell them
    apart: the stored tab index is what keeps a render on the ingested tab."""
    path = build_case(tmp_path, "hidden_first_same_headers")
    parsed = _parse(path)

    assert parsed["sheet_name"] == "Current"
    assert parsed["read_options"] == {
        "version": 2, "sheet_index": 1, "header_row": 0, "zero_padded_columns": [],
    }
    assert _read(path, parsed) == [{"name": "Bob", "amount": 900}]
    assert parsed["sample_rows"] == [{"name": "Bob", "amount": 900}]
    # The same columns without read_options are a 0.4.0 dataset: first tab.
    assert list(iter_rows(path, parsed["columns"])) == [{"name": "Ann", "amount": 100}]


def test_every_tab_hidden_but_one_later(tmp_path):
    path = build_case(tmp_path, "very_hidden_first")
    parsed = _parse(path)
    assert parsed["sheet_name"] == "B"
    assert parsed["read_options"]["sheet_index"] == 1
    assert _read(path, parsed) == [{"y": 2}]


@pytest.mark.parametrize("name", ["csv_blank_header_then_data", "blank_first_row_then_data"])
def test_the_header_is_the_first_row_even_when_it_is_blank(tmp_path, name):
    """Skipping a blank first row took the first data row as the header: in
    ",\\nAnn,5\\nBob,6" Ann stopped being a document. The first row is the
    header, as in 0.4.0; a service decides whether a blank one is acceptable."""
    path = build_case(tmp_path, name)
    parsed = _parse(path)

    assert [c["original"] for c in parsed["columns"]] == ["", ""]
    assert parsed["read_options"]["header_row"] == 0
    assert parsed["row_count"] == 2
    assert _read(path, parsed) == [
        {"col_": "Ann", "col__2": 5}, {"col_": "Bob", "col__2": 6},
    ]


def test_read_options_header_row_is_honoured(tmp_path):
    path = tmp_path / "offset.csv"
    path.write_text("title,\n,\nName,Amount\nAnn,5\n")
    columns = [{"original": "Name", "key": "name"}, {"original": "Amount", "key": "amount"}]
    rows = iter_rows(str(path), columns, read_options={"version": 2, "header_row": 2})
    assert list(rows) == [{"name": "Ann", "amount": 5}]


@pytest.mark.parametrize(
    "read_options",
    [
        {"version": 3},
        {"version": 2, "sheet_index": 5},
        {"version": 2, "header_row": -1},
        {"version": 2, "sheet_index": "0"},
        {"version": 2, "zero_padded_columns": ["a"]},
    ],
)
def test_unreadable_read_options_fail_loudly(tmp_path, read_options):
    """A dataset written by a newer reader, or options that do not fit the
    file, must not quietly fall back to some other tab or row."""
    path = build_case(tmp_path, "two_visible_sheets")
    with pytest.raises(ValueError):
        list(iter_rows(path, LEGACY["two_visible_sheets"]["columns"], read_options=read_options))


def test_zero_padded_text_only_where_the_padding_shows(tmp_path):
    """A column becomes text only when its format (three or more zero
    placeholders, optionally split by - or space) makes Excel show something
    other than the bare number for at least one of its cells. ZIP codes keep
    their leading zeros; a quantity under "00" stays a number."""
    path = build_case(tmp_path, "zero_padded_formats")
    parsed = _parse(path)

    assert parsed["read_options"]["zero_padded_columns"] == [1, 3, 4, 6, 7]
    rows = _read(path, parsed)
    assert typed(rows) == [
        _t(name=("str", "Ann"), zip=("str", "02134"), qty=("int", 12), zip_4=("str", "02134-0001"),
           ssn=("str", "012-34-5678"), literal=("int", 5), wide=("str", "1234567"),
           neg=("str", "-00005"), price=("int", 5)),
        _t(name=("str", "Bob"), zip=("str", "90210"), qty=("int", 7), zip_4=("str", "90210-1234"),
           ssn=("str", "123-45-6789"), literal=("int", 6), wide=("str", "00012"),
           neg=("str", "00003"), price=("int", 6)),
    ]
    assert parsed["sample_rows"] == rows
    types = {c["key"]: c["inferred_type"] for c in parsed["columns"]}
    assert types == {
        "name": "string", "zip": "string", "qty": "number", "zip_4": "string", "ssn": "string",
        "literal": "number", "wide": "string", "neg": "string", "price": "number",
    }


def test_a_column_the_padding_never_changes_keeps_its_numbers(tmp_path):
    """int 12 under "00" used to become "12": {{ quantity * 2 }} went from 24
    to "1212". Under "000" with every value three digits wide, nothing is
    padded either, so the column stays numeric."""
    path = tmp_path / "qty.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Qty", "Code"])
    for qty, code in ((12, 100), (40, 999)):
        sheet.append([qty, code])
    for row in (2, 3):
        sheet[f"A{row}"].number_format = "00"
        sheet[f"B{row}"].number_format = "000"
    workbook.save(path)

    parsed = _parse(str(path))
    assert parsed["read_options"]["zero_padded_columns"] == []
    assert typed(_read(str(path), parsed)) == [
        _t(qty=("int", 12), code=("int", 100)), _t(qty=("int", 40), code=("int", 999)),
    ]


@pytest.mark.parametrize(
    "number_format",
    [
        "0" * 100_000,
        "0" * 21,
        "0" * 10 + "-" + "0" * 11,
        "000" + "-0" * 40,
        '0"0"', '00"0"', '0\\0', '00000"-"0000"0"',
    ],
)
def test_zero_pad_formats_that_are_too_long_or_hold_literal_zeros_are_ignored(tmp_path, number_format):
    """The format comes from the uploaded file. A 100,000-zero format turned
    every small integer into a 100 KB string before any cell limit applied,
    and a quoted literal zero (0"0" shows 5 as "50") was read as padding."""
    import time

    path = tmp_path / "crafted.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Code"])
    for value in range(1, 200):
        sheet.append([value])
        sheet.cell(row=value + 1, column=1).number_format = number_format
    workbook.save(path)

    started = time.perf_counter()
    parsed = _parse(str(path))
    rows = _read(str(path), parsed)
    assert time.perf_counter() - started < 5
    assert parsed["read_options"]["zero_padded_columns"] == []
    assert rows[:2] == [{"code": 1}, {"code": 2}]


def test_zero_pad_formats_at_the_limits_still_apply(tmp_path):
    path = tmp_path / "limits.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Twenty", "Quoted", "Escaped"])
    sheet.append([7, 21340001, 12345678])
    sheet["A2"].number_format = "0" * 20
    sheet["B2"].number_format = '00000"-"0000'
    sheet["C2"].number_format = "000\\ 00\\ 0000"
    workbook.save(path)

    parsed = _parse(str(path))
    assert _read(str(path), parsed) == [
        {"twenty": "0" * 19 + "7", "quoted": "02134-0001", "escaped": "012 34 5678"},
    ]


def test_grouping_and_the_cli_read_the_ingested_tab(tmp_path):
    """Grouping and the CLI read through iter_rows too; without read_options
    they would group the hidden tab's rows under the visible tab's columns."""
    import argparse

    from sheetrender.cli import _batch_contexts
    from sheetrender.grouping import (
        detect_group_candidates,
        grouped_render_units,
        scan_groups,
    )

    path = build_case(tmp_path, "hidden_first_same_headers")
    parsed = _parse(path)
    columns, options = parsed["columns"], parsed["read_options"]
    config = {"group_by": ["name"]}

    groups = grouped_render_units(path, columns, None, config, read_options=options)
    assert [group.rows for group in groups] == [[{"name": "Bob", "amount": 900}]]
    assert scan_groups(path, columns, config, read_options=options).sample_group["name"] == "Bob"
    detection = detect_group_candidates(path, columns, 1, read_options=options)
    assert detection["row_count"] == 1

    args = argparse.Namespace(group="none", group_by=None, data=path)
    assert _batch_contexts(args, parsed) == ([(0, {"name": "Bob", "amount": 900})], False)
    args = argparse.Namespace(group="none", group_by="name", data=path)
    contexts, grouped = _batch_contexts(args, parsed)
    assert grouped and contexts[0][1]["name"] == "Bob"


def test_the_cli_refuses_a_file_whose_first_row_is_blank(tmp_path):
    """The SheetRender app refuses these uploads with the same advice; the
    first row is always the header, so a blank one names no columns. batch
    refuses them; inspect still shows the file as it is."""
    from sheetrender.cli import CliError, _parse_dataset

    for name in ("csv_blank_header_then_data", "table_starts_lower_down", "csv_blank_lines_above_header"):
        with pytest.raises(CliError, match="first row"):
            _parse_dataset(build_case(tmp_path, name), require_header=True)


def test_parse_reads_the_tab_and_header_row_a_caller_chose(tmp_path):
    """A service that lets the user pick the tab or the header row (a title
    row above the headings) parses with that choice, and the read_options it
    gets back make iter_rows read the same way."""
    from sheetrender.sheets import parse_csv, parse_xlsx

    path = tmp_path / "choice.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "Summary"
    first.append(["Total", 5])
    second = workbook.create_sheet("Orders")
    second.append(["Quarterly orders"])
    second.append([])
    second.append(["Name", "Zip"])
    second.append(["Ann", 2134])
    second["B4"].number_format = "00000"
    workbook.save(path)

    parsed = parse_xlsx(str(path), sheet_index=1, header_row=2)
    assert parsed["sheet_name"] == "Orders"
    assert parsed["read_options"] == {
        "version": 2, "sheet_index": 1, "header_row": 2, "zero_padded_columns": [1],
    }
    assert [c["original"] for c in parsed["columns"]] == ["Name", "Zip"]
    assert parsed["sample_rows"] == _read(str(path), parsed) == [{"name": "Ann", "zip": "02134"}]

    csv_path = tmp_path / "choice.csv"
    csv_path.write_text("Quarterly orders\n\nName,Amount\nAnn,5\n")
    parsed = parse_csv(str(csv_path), header_row=2)
    assert parsed["read_options"] == {"version": 2, "header_row": 2}
    assert parsed["sample_rows"] == _read(str(csv_path), parsed) == [{"name": "Ann", "amount": 5}]

    with pytest.raises(ValueError):
        parse_xlsx(str(path), sheet_index=2)
    with pytest.raises(ValueError):
        parse_xlsx(str(path), header_row=-1)


def test_read_options_with_keys_this_version_does_not_know_are_refused(tmp_path):
    """A later reader may add a key (a delimiter, say). Ignoring it would read
    the file some other way, so an older reader refuses instead."""
    path = build_case(tmp_path, "csv_trailing_commas")
    with pytest.raises(ValueError, match="delimiter"):
        list(iter_rows(path, [{"key": "a"}], read_options={"version": 2, "delimiter": "\t"}))
